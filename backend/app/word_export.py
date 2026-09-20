"""正式工单 Word 导出。

只依据服务端已保存的 IssueRecord 生成；浏览器不能提交工单正文、编号、风险或状态。
没有的数据写"待补充"或"无"，不出现内部字段名、JSON、Schema 等开发内容。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PENDING = "待补充"
EMPTY = "无"

CATEGORY_LABELS = {
    "safety": "安全", "quality": "质量", "management": "管理",
    "logistics": "后勤", "consultation": "咨询", "unknown": "待判断",
}
RISK_LABELS = {
    "undetermined": "待判断", "low": "低", "medium": "中",
    "high": "高", "emergency": "紧急",
}
STATUS_LABELS = {
    "draft": "草稿", "submitted": "已提交", "assigned": "已派工",
    "rectifying": "整改中", "pending_review": "待复查", "closed": "已关闭",
}
ROLE_LABELS = {
    "safety_officer": "安全管理人员", "quality_inspector": "质量检查人员",
    "site_manager": "现场管理人员", "facilities_staff": "后勤维修人员",
}
AI_NOTICE = "AI辅助信息仅供现场人员参考，不构成最终专业认定。"


def _text(value: Any, *, empty: str = PENDING) -> str:
    if value is None:
        return empty
    text = str(value).strip()
    return text or empty


def _joined(values: Iterable[Any] | None, *, empty: str = EMPTY) -> str:
    items = [str(value).strip() for value in (values or []) if str(value).strip()]
    return "；".join(items) or empty


def _label(value: Any, labels: Mapping[str, str], *, empty: str = PENDING) -> str:
    key = str(value)
    return labels.get(key, _text(value, empty=empty))


def _add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(20)


def _add_section(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12.5)


def _add_rule(document: Document) -> None:
    """一条浅色分隔线，让版式更接近正式纸质表单。"""

    paragraph = document.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "A6A6A6")
    borders.append(bottom)
    properties.append(borders)


def _add_table(document: Document, rows: Iterable[tuple[str, str]]) -> None:
    items = list(rows)
    table = document.add_table(rows=len(items), cols=2)
    table.style = "Table Grid"
    table.autofit = True
    for index, (label, value) in enumerate(items):
        cells = table.rows[index].cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)


def _event_time(record: Any, action: str) -> str | None:
    for event in getattr(record, "events", ()) or ():
        if getattr(event, "action", None) == action:
            return _text(getattr(event, "occurred_at", None), empty=PENDING)
    return None


def build_issue_record_docx(
    record: Any,
    linked_photos: Iterable[Mapping[str, Any]] = (),
) -> bytes:
    """把一条正式 IssueRecord 渲染成可打印、可归档的 .docx。"""

    # 本轮正式工单模板不再单独列出图片证据；参数保留以兼容既有调用方。
    _ = list(linked_photos)
    review = record.review_fields
    analysis = record.analysis

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    _add_title(document, "施工现场问题工单")
    _add_rule(document)

    _add_section(document, "一、基本信息")
    _add_table(document, [
        ("工单编号", _text(record.record_id)),
        ("生成时间", _text(record.created_at)),
        ("问题类别", _label(analysis.category, CATEGORY_LABELS)),
        ("风险等级", _label(analysis.risk_level, RISK_LABELS)),
        ("项目", _text(review.project)),
        ("区域", _text(review.area)),
    ])
    _add_rule(document)

    _add_section(document, "二、问题情况")
    _add_table(document, [
        ("问题标题", _text(review.record_title)),
        ("问题描述", _text(review.record_description)),
    ])
    _add_rule(document)

    _add_section(document, "三、AI辅助建议")
    _add_table(document, [
        ("AI简要建议", _joined(analysis.suggested_actions)),
        ("是否需要人工复核", "需要" if analysis.requires_human_review else "不需要"),
    ])
    document.add_paragraph(AI_NOTICE)
    _add_rule(document)

    _add_section(document, "四、处理情况")
    _add_table(document, [
        ("当前状态", "已取消" if str(record.disposition) == "cancelled"
         else _label(record.status, STATUS_LABELS)),
        ("责任角色建议", _label(record.suggested_responsible_role, ROLE_LABELS)),
        ("实际负责人", _text(record.assigned_to)),
        ("处理说明", _text(record.information_request or record.rectification_note)),
        ("整改结果", _text(record.review_note)),
    ])
    _add_rule(document)

    _add_section(document, "五、记录信息")
    _add_table(document, [
        ("提交时间", _event_time(record, "submit") or PENDING),
        ("最近更新时间", _text(record.updated_at)),
        ("关闭时间", _event_time(record, "close") or PENDING),
    ])

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_filename(record_id: str) -> str:
    return f"{record_id}_施工现场问题工单.docx"


def export_headers(record_id: str, filename: str) -> dict[str, str]:
    from urllib.parse import quote

    return {
        "Content-Disposition": (
            f"attachment; filename=\"{record_id}.docx\"; "
            f"filename*=UTF-8''{quote(filename)}"
        )
    }
