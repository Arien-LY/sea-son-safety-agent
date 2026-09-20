"""正式工单 Word 导出：只依据服务端已保存的 record_id。"""

from io import BytesIO

from docx import Document

from test_phase3_api import confirmed_payload, phase3_client

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def create_record(client, suffix="word-export-001", category="safety", risk_level="medium"):
    confirmation = confirmed_payload(client, category=category, risk_level=risk_level, suffix=suffix)
    response = client.post("/api/issue-records", json={
        "confirmation": confirmation, "idempotency_key": f"create-{suffix}",
        "actor": {"actor_id": "reporter-word", "role": "reporter"}})
    assert response.status_code == 200, response.text
    return response.json()["record"]


def test_saved_record_exports_a_reopenable_docx(tmp_path):
    client = phase3_client(tmp_path)
    record = create_record(client)
    response = client.get(f"/api/issue-records/{record['record_id']}/export.docx")
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition and record["record_id"] in disposition
    assert "%E6%96%BD%E5%B7%A5" in disposition  # 文件名包含“施工”
    document = Document(BytesIO(response.content))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    body += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for token in ["施工现场问题工单", "一、基本信息", "二、问题情况", "三、AI辅助建议",
                  "四、处理情况", "五、记录信息", "工单编号", record["record_id"],
                  "问题类别", "风险等级", "问题标题", "问题描述", "当前状态", "待补充",
                  "AI辅助信息仅供现场人员参考，不构成最终专业认定。"]:
        assert token in body, token
    # 正式模板不再单独列出图片证据，也不出现内部开发字段名。
    for forbidden in ["现场证据", "IssueAnalysis", "recommended_route", "confidence",
                      "hmac", "revision", "schema"]:
        assert forbidden not in body, forbidden


def test_missing_record_id_does_not_generate_a_file(tmp_path):
    client = phase3_client(tmp_path)
    response = client.get("/api/issue-records/ISS-000000000000/export.docx")
    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/json")


def test_unsaved_proposal_cannot_be_exported(tmp_path):
    client = phase3_client(tmp_path)
    confirmed_payload(client, category="safety", risk_level="medium", suffix="word-unsaved")
    assert client.get("/api/issue-records").json()["total"] == 0
    assert client.get("/api/issue-records/ISS-000000000000/export.docx").status_code == 404
